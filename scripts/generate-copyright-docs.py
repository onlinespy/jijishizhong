from pathlib import Path
import json
import textwrap

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "generated"
SOFTWARE_NAME = "记记食钟健康管理软件"
VERSION = "V1.0"


def set_run_font(run, name="宋体", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_paragraph_spacing(paragraph, before=0, after=0, line=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if line is not None:
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.line_spacing = Pt(line)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)
    set_run_font(run, size=8, color=(90, 90, 90))


def configure_document(doc, header_title):
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.text = header_title
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in paragraph.runs:
        set_run_font(run, size=8.5, color=(80, 80, 80))
    add_page_number(paragraph)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_paragraph.add_run(f"{SOFTWARE_NAME} {VERSION}")
    set_run_font(footer_run, size=8, color=(110, 110, 110))


def add_title(doc, title, subtitle=""):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(paragraph, before=100, after=18, line=30)
    run = paragraph.add_run(title)
    set_run_font(run, name="黑体", size=24, bold=True)

    if subtitle:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(paragraph, after=40, line=24)
        run = paragraph.add_run(subtitle)
        set_run_font(run, size=14)


def add_info_line(doc, label, value):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(paragraph, after=8, line=18)
    run = paragraph.add_run(f"{label}：{value}")
    set_run_font(run, size=11)


def add_heading(doc, text, level=1, page_break=False):
    if page_break and doc.paragraphs:
        doc.add_page_break()
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, before=10 if level == 1 else 7, after=6, line=20)
    run = paragraph.add_run(text)
    set_run_font(run, name="黑体", size=16 if level == 1 else 13, bold=True)
    return paragraph


def add_body(doc, text, indent=False):
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, after=4, line=17)
    if indent:
        paragraph.paragraph_format.first_line_indent = Pt(21)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph()
        set_paragraph_spacing(paragraph, after=3, line=16)
        paragraph.paragraph_format.left_indent = Pt(16)
        run = paragraph.add_run(f"• {item}")
        set_run_font(run, size=10.5)


def add_screenshot_placeholder(doc, caption, rows=10):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    for index in range(rows):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(paragraph, before=0, after=0, line=16)
        text = f"【截图占位：{caption}】" if index == 3 else ""
        run = paragraph.add_run(text)
        set_run_font(run, size=10, color=(120, 120, 120))
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(caption_paragraph, before=4, after=10, line=16)
    run = caption_paragraph.add_run(f"图：{caption}")
    set_run_font(run, size=9, color=(90, 90, 90))


def ordered_source_files():
    files = [
        ROOT / "miniprogram" / "app.js",
        ROOT / "miniprogram" / "app.json",
        ROOT / "miniprogram" / "app.wxss",
        ROOT / "miniprogram" / "sitemap.json",
        ROOT / "miniprogram" / "config" / "index.js",
    ]
    files.extend([
        ROOT / "miniprogram" / "custom-tab-bar" / "index.js",
        ROOT / "miniprogram" / "custom-tab-bar" / "index.json",
        ROOT / "miniprogram" / "custom-tab-bar" / "index.wxml",
        ROOT / "miniprogram" / "custom-tab-bar" / "index.wxss",
    ])
    files.extend(sorted((ROOT / "miniprogram" / "utils").glob("*.js")))
    files.extend(sorted((ROOT / "miniprogram" / "data").glob("*.js")))

    app_json = json.loads((ROOT / "miniprogram" / "app.json").read_text(encoding="utf-8"))
    for page in app_json["pages"]:
        base = ROOT / "miniprogram" / page
        for extension in (".js", ".wxml", ".wxss", ".json"):
            target = Path(str(base) + extension)
            if target.exists():
                files.append(target)

    for function_name in ("login", "health", "sendReminders"):
        base = ROOT / "cloudfunctions" / function_name
        files.append(base / "index.js")
        files.append(base / "package.json")
        config_file = base / "config.json"
        if config_file.exists():
            files.append(config_file)

    return files


def collect_source_lines():
    lines = []
    for file_path in ordered_source_files():
        relative = file_path.relative_to(ROOT).as_posix()
        lines.append(f"===== FILE: {relative} =====")
        for raw_line in file_path.read_text(encoding="utf-8").splitlines():
            line = raw_line.replace("\t", "  ").rstrip()
            if not line:
                continue
            wrapped = textwrap.wrap(
                line,
                width=108,
                replace_whitespace=False,
                drop_whitespace=False,
                break_long_words=True,
                break_on_hyphens=False,
            )
            lines.extend(wrapped or [""])
    return lines


def build_source_document():
    doc = Document()
    configure_document(doc, f"{SOFTWARE_NAME} {VERSION}  源程序")

    add_title(doc, f"{SOFTWARE_NAME} {VERSION}", "源程序鉴别材料")
    add_info_line(doc, "著作权人", "请填写个人真实姓名或登记主体名称")
    add_info_line(doc, "开发完成日期", "请填写实际完成日期")
    add_info_line(doc, "提交方式", "前 30 页 + 后 30 页，每页 50 行")
    add_info_line(doc, "材料生成日期", "2026-10-10")
    add_body(
        doc,
        "说明：本文件按常见软著材料格式生成，共包含前 30 页和后 30 页源程序草案。"
        "正式提交前应根据受理机构当期要求核对页眉、页码、署名、字体和份数。",
        indent=True,
    )
    doc.add_page_break()

    all_lines = collect_source_lines()
    page_size = 50
    first_lines = all_lines[: page_size * 30]
    last_lines = all_lines[-(page_size * 30):]

    code_style = doc.styles.add_style("CodeLine", WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    code_style.font.size = Pt(8.2)
    code_style.paragraph_format.space_before = Pt(0)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    code_style.paragraph_format.line_spacing = Pt(10.5)

    def render_pages(lines, section_label):
        for page_index in range(0, len(lines), page_size):
            chunk = lines[page_index : page_index + page_size]
            if page_index > 0:
                doc.add_page_break()
            for line in chunk:
                paragraph = doc.add_paragraph(style="CodeLine")
                paragraph.paragraph_format.keep_together = True
                run = paragraph.add_run(line)
                set_run_font(run, name="Consolas", size=8.2)

    render_pages(first_lines, "第一部分：源程序前 30 页")
    render_pages(last_lines, "第二部分：源程序后 30 页")

    output = OUTPUT / f"{SOFTWARE_NAME}{VERSION}-源程序前后60页草案.docx"
    doc.save(output)
    return output


MANUAL_SECTIONS = [
    {
        "title": "一、软件概述",
        "paragraphs": [
            "记记食钟是一款面向普通成年人的健康生活记录类微信小程序，核心功能包括轻断食计划打卡、饮食热量记录、运动消耗记录、体重追踪、喝水打卡、经期记录和久坐体态科普。",
            "本软件仅进行数据记录、简单计算与统计展示，不提供医疗诊断、诊疗建议或疗效承诺，不设置社交社区、商城和积分兑换功能。",
            "软件采用微信原生小程序界面和微信云开发数据能力，用户无需安装独立应用，可通过微信直接完成日常记录。",
        ],
        "screenshots": ["首页今日仪表盘总览"],
    },
    {
        "title": "二、运行环境与技术架构",
        "paragraphs": [
            "客户端运行于微信 iOS 和 Android 平台，页面采用 WXML、WXSS 和 JavaScript 开发。数据层支持本地存储与微信云开发两种模式。",
            "云开发模式通过 health 云函数读写 health_profiles 集合，以 OpenID 隔离用户数据。login 云函数用于获取当前用户标识，sendReminders 云函数用于处理用户授权后的订阅消息。",
            "在未部署云函数时，软件可使用本地模式运行，便于开发调试和功能演示；正式使用时建议启用云开发，以便在用户更换设备后继续读取自己的数据。",
            "界面以 375px 手机竖屏为统一设计基准，页面左右固定 16px 内边距，使用深绿色、浅绿色和中性灰组成的固定配色，卡片圆角 12px、间距 16px、主按钮高度 44px，所有页面复用同一套组件规范。",
        ],
        "screenshots": ["云开发数据与用户隔离示意"],
    },
    {
        "title": "三、导航结构",
        "paragraphs": [
            "底部导航包含首页、记录、科普、我的四项。首页用于查看今日概览；记录用于录入饮食、运动、体重和经期；科普用于阅读日常体态内容；我的用于管理档案、日历、勋章、设置和协议。",
            "底部导航采用自定义组件，使用 Lucide 线性图标、加大文字、选中底色和顶部指示条，提升入口辨识度。页面上方根据页面功能显示不同标题。首页支持下拉同步云端数据，其他页面在返回时自动刷新当前记录。",
        ],
        "screenshots": ["底部四项导航结构"],
    },
    {
        "title": "四、首次使用与健康风险确认",
        "paragraphs": [
            "用户首次进入断食计划功能时，页面会显示健康风险提示，内容包含孕妇、哺乳期、低血糖、胃病、糖尿病和青少年不建议进行轻断食的说明。",
            "用户必须手动勾选“我已阅读并理解以上提示”，确认按钮才会变为可用状态。确认后设定值写入用户数据，后续进入不再重复强制弹窗。",
        ],
        "screenshots": ["首次进入断食功能的健康风险确认弹窗"],
    },
    {
        "title": "五、轻断食计划设置",
        "paragraphs": [
            "软件提供 16+8、14+10、18+6 和 5+2 四种预设模板。选择每日时间窗模板后，用户可修改进食开始时间和结束时间，页面会同步计算进食与断食小时数。",
            "5+2 模板采用每周计划记录方式，用户自行选择两天进行计划内打卡，软件不安排代餐、食谱和热量目标。保存计划后，首页显示时间轴、当前窗口状态和实时倒计时。",
            "每日完成后，用户可点击“记录今日完成”。打卡记录写入日历，达到连续记录条件后自动展示相应勋章。",
            "首页和计划设置页均提供“停止当前轻断食计划”入口。停止后只关闭倒计时和每日打卡入口，历史打卡与计划设置保留，用户可随时重新启用。",
        ],
        "screenshots": ["轻断食模板与进食窗口设置", "首页断食时间轴、倒计时和停止计划入口"],
    },
    {
        "title": "六、饮食记录",
        "paragraphs": [
            "用户先在早餐、午餐、晚餐、加餐中选择餐次，再选择基础食物库或手动录入。食物库按照主食、肉蛋奶、蔬果、家常菜分组。",
            "选择食物后，用户填写克数或碗数，系统按内置每 100 克参考热量计算并回显参考热量。对于没有纳入食物库的内容，可手动填写名称、分量、单位和参考热量。",
            "保存后的记录进入当日列表，首页自动汇总当日摄入总热量。软件只展示数字，不对饮食内容作好坏评价。",
        ],
        "screenshots": ["基础食物库与食用量换算", "当日饮食记录和热量汇总"],
    },
    {
        "title": "七、运动记录",
        "paragraphs": [
            "基础运动库包含快走、慢跑、跳绳、瑜伽、健身操、骑行、游泳、爬楼梯。用户选择运动项目并填写分钟数，系统读取档案体重后按 MET 公式估算运动消耗。",
            "如基础运动库没有目标项目，用户可选择自定义运动，填写名称、时长和参考消耗热量。首页将展示摄入热量、运动消耗和二者差额。",
            "运动消耗为估算值，受强度、效率、体重和个体状态影响，页面已提示仅作记录参考。",
        ],
        "screenshots": ["基础运动库与运动时长录入", "当日运动记录和消耗汇总"],
    },
    {
        "title": "八、体重与 BMI",
        "paragraphs": [
            "用户在记录页填写体重并保存。同一天重复保存时会更新原记录，不会形成重复日期。软件读取档案身高计算 BMI，并只展示数值。",
            "历史体重记录按日期倒序排列，最近最多 14 次记录使用 Canvas 绘制简易折线图。页面建议用户每周记录 1 至 3 次，不要求每天称重。",
        ],
        "screenshots": ["体重录入、BMI 数值与趋势折线图"],
    },
    {
        "title": "九、喝水记录与提醒",
        "paragraphs": [
            "每日喝水目标限制在 1500 至 2500ml。用户点击 200ml、300ml 或 500ml 快捷按钮即可记录，首页环形进度和当日总量实时更新。误触后可通过首页“撤销上一次”回退最近一次记录，也可在历史页逐条删除。",
            "提醒功能使用微信订阅消息。用户可以在首页或设置页主动授权，授权后才能由云函数发送提醒。用户关闭开关后不再发送。小程序不能后台常驻，也不绕过微信授权规则。",
        ],
        "screenshots": ["首页喝水环形进度与快捷打卡", "订阅消息授权设置"],
    },
    {
        "title": "十、经期记录",
        "paragraphs": [
            "用户选择经期开始和结束日期后保存。软件根据历史开始日期计算平均周期，并简单估算下次经期、排卵日和易孕期。",
            "所有结果均标注为预估参考，不能用于避孕、备孕诊断或其他医疗决策。用户可一键隐藏经期模块，隐藏后记录页不再显示相关入口，并可在档案页恢复。",
        ],
        "screenshots": ["经期日期录入与周期预估", "档案中的经期模块隐藏开关"],
    },
    {
        "title": "十一、体态科普",
        "paragraphs": [
            "科普页包含久坐办公体态、站姿习惯和居家简易拉伸三个栏目。文章内容聚焦日常习惯调整，不出现治疗、根治等医疗疗效描述。",
            "每篇文章使用独立详情页展示标题、栏目、阅读时间和分段内容。文章末尾统一附带日常健康科普免责声明。",
        ],
        "screenshots": ["科普栏目与文章列表", "科普文章详情与免责声明"],
    },
    {
        "title": "十二、个人健康档案",
        "paragraphs": [
            "用户可填写称呼、性别、生日、身高和体重，并可设置是否隐藏经期模块。身高用于 BMI 数值计算，体重用于运动消耗估算。",
            "档案支持随时修改。页面只显示身高、体重和 BMI 数值，不输出偏瘦、肥胖等体型判定文字。",
        ],
        "screenshots": ["个人健康档案编辑页"],
    },
    {
        "title": "十三、综合打卡日历",
        "paragraphs": [
            "综合日历使用颜色标记区分断食、喝水、体重和经期记录，同一天可同时显示多个标记。用户可切换月份并点击日期查看详情。",
            "详情区域展示当天断食打卡、喝水总量、体重、饮食摄入、运动消耗和经期区间。日历只回看用户录入的数据，不生成健康评分。",
        ],
        "screenshots": ["综合打卡日历月视图", "单个日期的综合记录详情"],
    },
    {
        "title": "十四、勋章墙",
        "paragraphs": [
            "勋章根据断食连续打卡、喝水连续打卡、体重记录次数和累计记录次数自动计算。已解锁和未解锁状态通过颜色与进度条区分。",
            "勋章仅作页面展示，不包含积分、等级、兑换或任何健康效果承诺。",
        ],
        "screenshots": ["勋章墙与解锁进度"],
    },
    {
        "title": "十五、设置与数据清理",
        "paragraphs": [
            "设置页可管理喝水提醒和经期提醒的订阅状态，并进入隐私政策和全局免责声明。",
            "清除数据功能要求二次确认。确认后清除断食、饮食、运动、体重、喝水和经期记录，个人档案、断食计划和提醒设置保留。启用云开发时同步清理云端记录。",
        ],
        "screenshots": ["设置页与订阅消息管理", "清除全部记录的二次确认"],
    },
    {
        "title": "十六、隐私与免责声明",
        "paragraphs": [
            "隐私政策说明软件收集的健康档案、打卡记录、订阅设置和云开发用户标识，并说明数据用途、存储方式、共享限制、删除方式和未成年人使用提示。",
            "全局免责声明说明软件属于健康记录辅助工具，不提供医疗诊断和诊疗建议；热量、运动消耗和经期推算仅为估算，不作为医学依据。",
        ],
        "screenshots": ["隐私政策页面", "全局免责声明页面"],
    },
    {
        "title": "十七、异常与边界处理",
        "paragraphs": [
            "未填写体重时，运动库估算会提示先完善档案；未选择食物或填写无效数量时，饮食记录不会保存；结束日期早于开始日期时，经期记录不会保存。",
            "订阅模板未配置或用户拒绝授权时，提醒开关保持关闭，不会影响记录、统计、日历和科普功能。云端同步失败时，页面继续保留本地数据并给出提示。",
        ],
        "screenshots": ["未授权订阅消息时的配置提示"],
    },
    {
        "title": "十八、发布与申请软著准备",
        "paragraphs": [
            "正式发布前应将 AppID 替换为真实小程序 AppID，将演示数据开关关闭，删除演示记录，配置隐私保护指引，并选择“生活服务-健康养生工具”类目。",
            "申请软著时，应准备申请表、身份证明、源程序前后各 30 页和软件说明书。软件名称建议统一使用“记记食钟健康管理软件 V1.0”。",
        ],
        "screenshots": ["微信公众平台类目与版本信息"],
    },
]


def build_manual_document():
    doc = Document()
    configure_document(doc, f"{SOFTWARE_NAME} {VERSION}  软件说明书")

    add_title(doc, f"{SOFTWARE_NAME} {VERSION}", "软件说明书")
    add_info_line(doc, "软件简称", "记记食钟")
    add_info_line(doc, "著作权人", "请填写个人真实姓名或登记主体名称")
    add_info_line(doc, "开发完成日期", "请填写实际完成日期")
    add_info_line(doc, "文档版本", "V1.0")
    add_body(
        doc,
        "本说明书根据当前项目代码和页面结构生成，用于软著申请材料整理。"
        "正式提交时，请替换截图占位并核对著作权人、完成日期、页面名称和版本号。",
        indent=True,
    )

    add_heading(doc, "目录", page_break=True)
    for index, section in enumerate(MANUAL_SECTIONS, 1):
        paragraph = doc.add_paragraph()
        set_paragraph_spacing(paragraph, after=4, line=17)
        run = paragraph.add_run(section["title"])
        set_run_font(run, size=11)

    for section in MANUAL_SECTIONS:
        add_heading(doc, section["title"], page_break=True)
        for text in section["paragraphs"]:
            add_body(doc, text, indent=True)
        for caption in section["screenshots"]:
            add_screenshot_placeholder(doc, caption)

    output = OUTPUT / f"{SOFTWARE_NAME}{VERSION}-软件说明书草案.docx"
    doc.save(output)
    return output


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    manual = build_manual_document()
    source = build_source_document()
    print(manual)
    print(source)


if __name__ == "__main__":
    main()
